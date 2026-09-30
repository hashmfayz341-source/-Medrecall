"""Three realistic lecture decks (python-pptx) exported to PDF by LibreOffice."""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
IMG = os.path.join(HERE, "img")

def deck():
    p = Presentation()
    p.slide_width, p.slide_height = Inches(13.333), Inches(7.5)
    return p

def title_slide(p, title, sub):
    s = p.slides.add_slide(p.slide_layouts[0])
    s.shapes.title.text = title
    s.placeholders[1].text = sub
    return s

def bullets(p, title, items, size=24):
    """items: str or (str, level)."""
    s = p.slides.add_slide(p.slide_layouts[1])
    s.shapes.title.text = title
    body = s.placeholders[1]
    body.left, body.top, body.width, body.height = Inches(0.7), Inches(1.6), Inches(12), Inches(5.6)
    tf = body.text_frame
    first = True
    for it in items:
        text, lvl = (it, 0) if isinstance(it, str) else it
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        para.text = text
        para.level = lvl
        for r in para.runs: r.font.size = Pt(size if lvl == 0 else size - 4)
    return s

def blank(p, title):
    s = p.slides.add_slide(p.slide_layouts[5])  # title only
    s.shapes.title.text = title
    return s

def textbox(s, x, y, w, h, paras, size=20, bold_first=False):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, t in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.text = t
        para.space_after = Pt(10)
        for r in para.runs:
            r.font.size = Pt(size)
            if bold_first and i == 0: r.font.bold = True
    return tb

def table(s, x, y, w, rows, size=16, row_h=0.5):
    t = s.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows))).table
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            c = t.cell(i, j)
            c.text = cell
            for para in c.text_frame.paragraphs:
                for r in para.runs: r.font.size = Pt(size)
    return t

def picture(s, name, x, y, w=None, h=None):
    kw = {}
    if w: kw["width"] = Inches(w)
    if h: kw["height"] = Inches(h)
    return s.shapes.add_picture(os.path.join(IMG, name), Inches(x), Inches(y), **kw)

def caption(s, x, y, w, text, size=14):
    return textbox(s, x, y, w, 0.5, [text], size=size)

# ---------------------------------------------------------------- Cell injury
def cell_injury():
    p = deck()
    title_slide(p, "Cell Injury, Adaptation and Cell Death", "Pathology — Lecture 3\nDr. Sarah Mahmoud\nFaculty of Medicine, 2025–2026")
    bullets(p, "Learning objectives", [
        "By the end of this lecture you should be able to:",
        ("List the causes of cell injury", 1),
        ("Describe the morphology of reversible and irreversible injury", 1),
        ("Compare necrosis and apoptosis", 1),
    ])
    bullets(p, "Causes of cell injury", [
        "Oxygen deprivation (hypoxia) is the most common cause of cell injury.",
        "Physical agents: trauma, extremes of temperature, radiation and electric shock",
        "Chemical agents and drugs",
        "Infectious agents: viruses, bacteria, fungi and parasites",
        "Immunologic reactions",
        "Genetic derangements",
        "Nutritional imbalances",
    ])
    s = blank(p, "Hypoxia versus ischemia")
    textbox(s, 0.7, 1.5, 12, 2.4, [
        "Hypoxia is a deficiency of oxygen that causes cell injury by reducing aerobic oxidative respiration.",
        "Ischemia is a loss of blood supply to a tissue. It impairs the delivery of substrates for glycolysis as well as oxygen. Therefore, ischemia injures tissues faster than hypoxia alone.",
    ])
    table(s, 1.5, 4.2, 10, [["Feature", "Hypoxia", "Ischemia"], ["Oxygen delivery", "Reduced", "Reduced"],
                            ["Anaerobic glycolysis", "Continues", "Stops"], ["Speed of injury", "Slower", "Faster"]])
    s = blank(p, "Reversible cell injury")
    textbox(s, 0.7, 1.5, 12, 5.5, [
        "The two main morphologic correlates of reversible cell injury are cellular swelling and fatty change.",
        "Cellular swelling is the first manifestation of almost all forms of injury to cells. It results from failure of energy-dependent ion pumps in the plasma membrane.",
        "Fatty change is seen mainly in cells involved in lipid metabolism, such as hepatocytes and myocardial cells. It is manifested by the appearance of lipid vacuoles in the cytoplasm.",
    ])
    bullets(p, "Mechanisms of cell injury", [
        "ATP depletion causes failure of the Na+/K+ ATPase pump, leading to influx of sodium and water and cellular swelling.",
        "Decreased ATP causes a switch to anaerobic glycolysis, which leads to lactic acid accumulation and decreased intracellular pH.",
        "Influx of calcium activates phospholipases, proteases, endonucleases and ATPases.",
        "Mitochondrial damage leads to formation of the mitochondrial permeability transition pore.",
        "Oxygen-derived free radicals cause lipid peroxidation of membranes.",
    ], size=22)
    s = blank(p, "Think about it")
    textbox(s, 0.7, 1.8, 12, 4, [
        "Why does ischemia cause more rapid injury than hypoxia?",
        "What happens to the cell if the insult persists?",
    ], size=28)
    s = blank(p, "Irreversible injury")
    textbox(s, 0.7, 1.4, 12, 2.2, [
        "Two phenomena consistently characterize irreversibility: the inability to restore mitochondrial function and profound disturbances in membrane function.",
        "There are two patterns of cell death: necrosis and apoptosis.",
    ])
    textbox(s, 0.7, 3.8, 12, 3.2, [
        "Nuclear changes in necrosis",
        "Pyknosis: nuclear shrinkage and increased basophilia",
        "Karyorrhexis: fragmentation of the pyknotic nucleus",
        "Karyolysis: fading of basophilia due to DNase activity",
    ], size=20, bold_first=True)
    s = blank(p, "Patterns of tissue necrosis")
    table(s, 0.6, 1.5, 12.1, [["Pattern", "Typical location", "Key feature"],
                              ["Coagulative necrosis", "Infarcts of heart and kidney", "Preserved tissue architecture"],
                              ["Liquefactive necrosis", "Brain infarcts and abscesses", "Digestion of dead cells into a liquid mass"],
                              ["Caseous necrosis", "Tuberculosis", "Friable, cheese-like white material"],
                              ["Fat necrosis", "Pancreas and breast", "Chalky white calcium soap deposits"],
                              ["Fibrinoid necrosis", "Blood vessel walls", "Deposition of immune complexes and fibrin"]], size=16, row_h=0.7)
    s = blank(p, "Gangrenous necrosis")
    textbox(s, 0.7, 1.3, 12, 1.6, [
        "Dry gangrene is coagulative necrosis of a limb that has lost its blood supply. Wet gangrene occurs when bacterial infection is superimposed, producing liquefactive necrosis.",
    ], size=18)
    picture(s, "wet_gangrene.jpg", 1.2, 3.0, w=4.8)
    picture(s, "dry_gangrene.jpg", 7.2, 3.0, w=4.8)
    caption(s, 1.2, 6.7, 4.8, "Figure 1: Wet gangrene of the foot")
    caption(s, 7.2, 6.7, 4.8, "Figure 2: Dry gangrene of the toes")
    s = blank(p, "Apoptosis: two pathways")
    textbox(s, 0.6, 1.4, 5.9, 5.6, [
        "Intrinsic (mitochondrial) pathway",
        "The intrinsic pathway is triggered by loss of survival signals, DNA damage and misfolded proteins. It is regulated by the BCL-2 family of proteins. Release of cytochrome c from mitochondria activates caspase-9.",
    ], size=18, bold_first=True)
    textbox(s, 6.9, 1.4, 5.9, 5.6, [
        "Extrinsic (death receptor) pathway",
        "The extrinsic pathway is initiated by engagement of death receptors such as Fas and the TNF receptor. It activates caspase-8. Both pathways converge on executioner caspases 3 and 6.",
    ], size=18, bold_first=True)
    s = blank(p, "Necrosis versus apoptosis")
    table(s, 0.8, 1.5, 11.7, [["Feature", "Necrosis", "Apoptosis"],
                              ["Cell size", "Enlarged (swelling)", "Reduced (shrinkage)"],
                              ["Plasma membrane", "Disrupted", "Intact"],
                              ["Adjacent inflammation", "Frequent", "None"],
                              ["Physiologic role", "Invariably pathologic", "Often physiologic"]], size=18, row_h=0.8)
    bullets(p, "Intracellular accumulations and calcification", [
        "Lipofuscin is a yellow-brown “wear-and-tear” pigment associated with aging.",
        "Hemosiderin is a hemoglobin-derived pigment that accumulates in tissues with iron excess.",
        "Dystrophic calcification occurs in dead or dying tissues despite normal serum calcium levels.",
        "Metastatic calcification occurs in normal tissues whenever there is hypercalcemia.",
    ], size=22)
    bullets(p, "References", [
        "Kumar V, Abbas AK, Aster JC. Robbins and Cotran Pathologic Basis of Disease. 10th ed. Elsevier; 2020.",
        "Lecture slides are for educational use only.",
    ], size=18)
    title_slide(p, "Thank you", "Questions?")
    return p

# ---------------------------------------------------------------- Pharmacology
def autonomic():
    p = deck()
    title_slide(p, "Autonomic Pharmacology", "Cholinergic and Adrenergic Drugs\nPharmacology — Lecture 7\nDr. Omar Haddad")
    bullets(p, "Outline", ["Neurotransmission in the ANS", "Receptors", "Cholinergic drugs", "Adrenergic drugs", "Case discussion"])
    s = blank(p, "Neurotransmitters of the autonomic nervous system")
    textbox(s, 0.7, 1.4, 12, 5.6, [
        "Acetylcholine is the neurotransmitter of all preganglionic autonomic fibers and of postganglionic parasympathetic fibers. Norepinephrine is the neurotransmitter of most postganglionic sympathetic fibers. Sweat glands are innervated by sympathetic cholinergic fibers.",
        "Acetylcholine is synthesized from choline and acetyl-CoA by choline acetyltransferase. It is hydrolyzed in the synaptic cleft by acetylcholinesterase.",
    ])
    s = blank(p, "Autonomic receptors")
    table(s, 0.6, 1.4, 12.1, [["Receptor", "G protein", "Second messenger", "Major location"],
                              ["M1", "Gq", "Increased IP3 and DAG", "CNS and gastric parietal cells"],
                              ["M2", "Gi", "Decreased cAMP", "Heart"],
                              ["M3", "Gq", "Increased IP3 and DAG", "Smooth muscle and exocrine glands"],
                              ["Alpha-1", "Gq", "Increased IP3 and DAG", "Vascular smooth muscle"],
                              ["Alpha-2", "Gi", "Decreased cAMP", "Presynaptic nerve terminals"],
                              ["Beta-1", "Gs", "Increased cAMP", "Heart and juxtaglomerular cells"],
                              ["Beta-2", "Gs", "Increased cAMP", "Bronchial smooth muscle"]], size=16, row_h=0.62)
    s = blank(p, "Direct-acting cholinergic agonists")
    textbox(s, 0.6, 1.4, 5.9, 5.6, [
        "Bethanechol",
        "Bethanechol is a muscarinic agonist that is resistant to acetylcholinesterase. It is used to treat postoperative urinary retention.",
    ], size=20, bold_first=True)
    textbox(s, 6.9, 1.4, 5.9, 5.6, [
        "Pilocarpine",
        "Pilocarpine is a muscarinic agonist used to treat glaucoma. It causes miosis and contraction of the ciliary muscle.",
    ], size=20, bold_first=True)
    bullets(p, "Indirect-acting agonists (cholinesterase inhibitors)", [
        "Neostigmine is a reversible cholinesterase inhibitor that does not cross the blood–brain barrier.",
        ("It is used to reverse neuromuscular blockade and to treat myasthenia gravis.", 1),
        "Physostigmine is a tertiary amine that crosses the blood–brain barrier. It is used to treat anticholinergic toxicity.",
        "Organophosphates irreversibly inhibit acetylcholinesterase.",
        "Pralidoxime regenerates acetylcholinesterase if given before aging of the enzyme.",
    ], size=22)
    s = blank(p, "Organophosphate poisoning")
    textbox(s, 0.7, 1.4, 6, 5.6, [
        "Muscarinic signs of organophosphate poisoning",
        "Diarrhea", "Urination", "Miosis", "Bronchospasm", "Bradycardia", "Lacrimation", "Salivation",
    ], size=20, bold_first=True)
    textbox(s, 7.0, 1.4, 5.8, 5.6, [
        "Atropine is given to block the muscarinic effects of organophosphate poisoning.",
        "Pralidoxime is given to reverse the nicotinic effects, such as muscle weakness.",
    ], size=20)
    s = blank(p, "Case discussion")
    textbox(s, 0.7, 1.4, 12, 5.6, [
        "A 45-year-old farmer presents with pinpoint pupils, excessive salivation and bradycardia.",
        "What is the most likely cause?",
        "Which antidote would you give first?",
    ], size=24)
    s = blank(p, "Muscarinic antagonists")
    textbox(s, 0.6, 1.4, 5.9, 4.2, [
        "Atropine",
        "Atropine is a competitive antagonist at muscarinic receptors. It increases heart rate and causes mydriasis. It is used to treat symptomatic bradycardia.",
    ], size=19, bold_first=True)
    textbox(s, 6.9, 1.4, 5.9, 4.2, [
        "Ipratropium and scopolamine",
        "Ipratropium is an inhaled muscarinic antagonist used in chronic obstructive pulmonary disease. Scopolamine is used to prevent motion sickness.",
    ], size=19, bold_first=True)
    textbox(s, 0.6, 5.5, 12.2, 1.6, [
        "Adverse effects of muscarinic antagonists include dry mouth, blurred vision, urinary retention, constipation and hyperthermia.",
    ], size=18)
    s = blank(p, "Adrenergic agonists")
    table(s, 0.6, 1.4, 12.1, [["Drug", "Receptor selectivity", "Clinical use"],
                              ["Epinephrine", "Alpha and beta", "Anaphylaxis"],
                              ["Norepinephrine", "Alpha-1, alpha-2 and beta-1", "Septic shock"],
                              ["Dobutamine", "Beta-1 more than beta-2", "Acute heart failure"],
                              ["Phenylephrine", "Alpha-1", "Nasal decongestion"],
                              ["Albuterol", "Beta-2", "Acute asthma"],
                              ["Clonidine", "Alpha-2", "Hypertension"]], size=17, row_h=0.65)
    s = blank(p, "Epinephrine and clonidine")
    textbox(s, 0.6, 1.4, 5.9, 5.6, [
        "Epinephrine",
        "Epinephrine is the drug of choice for anaphylaxis. At low doses it has predominantly beta effects. At high doses alpha-1 effects predominate and cause vasoconstriction.",
    ], size=19, bold_first=True)
    textbox(s, 6.9, 1.4, 5.9, 5.6, [
        "Clonidine",
        "Clonidine is a centrally acting alpha-2 agonist. It decreases sympathetic outflow from the brainstem. Abrupt withdrawal of clonidine may cause rebound hypertension.",
    ], size=19, bold_first=True)
    bullets(p, "Indirect sympathomimetics", [
        "Amphetamine releases stored catecholamines from nerve terminals.",
        "Cocaine blocks the reuptake of norepinephrine, dopamine and serotonin.",
        "Cocaine is the only local anesthetic that causes vasoconstriction.",
    ])
    bullets(p, "Alpha blockers", [
        "Prazosin is a selective alpha-1 antagonist used to treat hypertension and benign prostatic hyperplasia. Its main adverse effect is first-dose orthostatic hypotension.",
        "Phenoxybenzamine is an irreversible nonselective alpha antagonist used to treat pheochromocytoma before surgery.",
    ], size=22)
    s = blank(p, "Beta blockers")
    table(s, 0.6, 1.4, 12.1, [["Drug", "Selectivity", "Additional use"],
                              ["Propranolol", "Nonselective (beta-1 and beta-2)", "Migraine prophylaxis"],
                              ["Metoprolol", "Beta-1 selective", "Heart failure"],
                              ["Timolol", "Nonselective", "Glaucoma"]], size=17, row_h=0.6)
    textbox(s, 0.6, 4.2, 12.1, 3, [
        "Beta blockers decrease heart rate and myocardial contractility.",
        "Nonselective beta blockers can precipitate bronchospasm in patients with asthma.",
        "Beta blockers may mask the symptoms of hypoglycemia in diabetic patients.",
    ], size=19)
    bullets(p, "Summary", ["Cholinergic agonists → parasympathetic effects", "Antagonists → the opposite", "Know the receptor table!"])
    bullets(p, "References", ["Katzung BG. Basic and Clinical Pharmacology. 15th ed. McGraw-Hill; 2021.",
                              "Whalen K. Lippincott Illustrated Reviews: Pharmacology. 7th ed."], size=18)
    return p

# ---------------------------------------------------------------- Image-heavy pathology
def image_review():
    p = deck()
    def slide(title):
        s = blank(p, title)
        picture(s, "logo.png", 12.3, 0.15, w=0.8)  # repeated branding
        return s
    s = title_slide(p, "Gross and Microscopic Pathology", "Image review session — Lecture 9\nDepartment of Pathology")
    picture(s, "logo.png", 6.0, 0.4, w=1.3)
    s = slide("Wet and dry gangrene")
    picture(s, "wet_gangrene_b.jpg", 1.0, 1.4, w=5.2)
    picture(s, "dry_gangrene_b.jpg", 7.0, 1.4, w=5.2)
    caption(s, 3.0, 5.4, 7.5, "Figure 3: Wet (left) and dry (right) gangrene")
    textbox(s, 0.8, 5.9, 11.8, 1.5, [
        "Wet gangrene shows a swollen, moist limb with a poorly defined line of demarcation. Dry gangrene shows a shrunken, black, mummified limb with a clear line of demarcation.",
    ], size=17)
    s = slide("Thyroid carcinoma")
    picture(s, "papillary_ptc.jpg", 1.0, 1.4, w=4.8)
    picture(s, "follicular_ftc.jpg", 7.2, 1.4, w=4.8)
    caption(s, 1.0, 5.1, 4.8, "Papillary thyroid carcinoma")
    caption(s, 7.2, 5.1, 4.8, "Follicular thyroid carcinoma")
    textbox(s, 0.8, 5.7, 11.8, 1.7, [
        "Papillary carcinoma shows psammoma bodies and nuclei with a ground-glass appearance. Follicular carcinoma is diagnosed by capsular or vascular invasion.",
    ], size=17)
    s = slide("Diabetes mellitus: the islets")
    picture(s, "t1dm_insulitis.jpg", 1.0, 1.4, w=4.8)
    picture(s, "t2dm_amyloid.jpg", 7.2, 1.4, w=4.8)
    caption(s, 1.0, 5.1, 4.8, "Type 1 diabetes: insulitis with lymphocytes")
    caption(s, 7.2, 5.1, 4.8, "Type 2 diabetes: amyloid deposition in an islet")
    textbox(s, 0.8, 5.7, 11.8, 1.7, [
        "Type 1 diabetes mellitus results from autoimmune destruction of pancreatic beta cells. Type 2 diabetes mellitus is characterized by insulin resistance and relative insulin deficiency.",
    ], size=17)
    s = slide("Normal and cirrhotic liver")
    picture(s, "normal_liver.jpg", 1.0, 1.4, w=4.8)
    picture(s, "cirrhotic_liver.jpg", 7.2, 1.4, w=4.8)
    caption(s, 1.0, 5.1, 4.8, "Normal liver")
    caption(s, 7.2, 5.1, 4.8, "Cirrhotic liver with regenerative nodules")
    textbox(s, 0.8, 5.7, 11.8, 1.7, [
        "Cirrhosis is characterized by bridging fibrous septa and regenerative parenchymal nodules.",
    ], size=17)
    s = slide("Heart failure")
    picture(s, "left_hf_lung.jpg", 1.0, 1.4, w=4.8)
    picture(s, "right_hf_nutmeg.jpg", 7.2, 1.4, w=4.8)
    caption(s, 1.0, 5.1, 4.8, "Left-sided heart failure: pulmonary edema")
    caption(s, 7.2, 5.1, 4.8, "Right-sided heart failure: nutmeg liver")
    textbox(s, 0.8, 5.7, 11.8, 1.7, [
        "Left-sided heart failure most commonly causes pulmonary congestion and edema. Right-sided heart failure is most often caused by left-sided heart failure. It causes congestive hepatomegaly with a nutmeg appearance.",
    ], size=16)
    s = slide("Tuberculous granuloma")
    picture(s, "tb_granuloma.jpg", 1.0, 1.4, w=6.0)
    textbox(s, 3.3, 3.0, 2.2, 0.5, ["Caseous necrosis"], size=14)
    textbox(s, 5.2, 1.7, 2.2, 0.5, ["Langhans giant cell"], size=14)
    caption(s, 1.0, 5.9, 6.0, "Figure 7: Tuberculous granuloma")
    textbox(s, 7.4, 1.6, 5.4, 5, [
        "A tuberculous granuloma consists of central caseous necrosis surrounded by epithelioid macrophages and Langhans giant cells.",
    ], size=18)
    s = slide("Amyloidosis")
    picture(s, "amyloid_table.png", 0.8, 1.4, w=6.0)
    caption(s, 0.8, 5.0, 6.0, "Table 2: Classification of amyloidosis")
    picture(s, "congo_red.jpg", 7.3, 1.4, w=5.0)
    caption(s, 7.3, 5.2, 5.0, "Congo red stain under polarized light")
    textbox(s, 0.8, 5.8, 11.8, 1.5, [
        "Amyloid shows apple-green birefringence under polarized light after Congo red staining.",
    ], size=17)
    s = slide("Fatty change")
    picture(s, "thumb.jpg", 11.8, 6.6, w=0.35)  # tiny thumbnail
    picture(s, "fatty_liver.jpg", 1.0, 1.4, w=5.5)
    caption(s, 1.0, 5.6, 5.5, "Figure 8: Hepatic steatosis")
    textbox(s, 7.0, 1.6, 5.8, 5, [
        "Hepatic steatosis is most often caused by alcohol abuse and nonalcoholic fatty liver disease. Lipid accumulates as clear vacuoles in the cytoplasm of hepatocytes.",
    ], size=18)
    s = slide("Myocardial infarction")
    picture(s, "acute_mi.jpg", 1.0, 1.4, w=5.8)
    caption(s, 1.0, 5.8, 5.8, "Figure 9: Myocardial infarction at 3 days with neutrophils")
    textbox(s, 7.2, 1.6, 5.6, 5, [
        "Neutrophil infiltration peaks 1 to 3 days after myocardial infarction.",
        "Macrophages remove dead myocytes 3 to 7 days after infarction.",
    ], size=18)
    s = slide("Acute and chronic inflammation")
    picture(s, "acute_appendicitis.jpg", 1.0, 1.4, w=4.8)
    picture(s, "chronic_cholecystitis.jpg", 7.2, 1.4, w=4.8)
    caption(s, 1.0, 5.1, 4.8, "Acute appendicitis with neutrophils")
    caption(s, 7.2, 5.1, 4.8, "Chronic cholecystitis with lymphocytes")
    textbox(s, 0.8, 5.7, 11.8, 1.7, [
        "Acute inflammation is dominated by neutrophils. Chronic inflammation is characterized by lymphocytes, plasma cells and macrophages.",
    ], size=17)
    title_slide(p, "End of image review", "Thank you")
    return p

if __name__ == "__main__":
    # Then: soffice --headless --convert-to pdf --outdir .. pptx/*.pptx
    out = os.path.join(HERE, "pptx"); os.makedirs(out, exist_ok=True)
    cell_injury().save(os.path.join(out, "cell-injury.pptx"))
    autonomic().save(os.path.join(out, "autonomic-pharmacology.pptx"))
    image_review().save(os.path.join(out, "pathology-images.pptx"))
    print("ok")
